import io
import json
import os
import shutil
import tempfile
import time
import zipfile
from pathlib import Path

os.environ.setdefault("DATA_DIR", tempfile.mkdtemp())

import pikepdf
import pymupdf
import pytest
from PIL import Image, ImageDraw

from app.ops import UserError, analyze, convert, pages, secure

needs = lambda tool: pytest.mark.skipif(not shutil.which(tool), reason=f"{tool} not installed")


def make_pdf(path, n=3, prefix="Halaman", rotation=0):
    doc = pymupdf.open()
    for i in range(n):
        page = doc.new_page()
        page.insert_text((72, 100), f"{prefix} {i + 1}", fontsize=14)
        page.insert_text((300, 500), "SECRET", fontsize=14)
        page.set_rotation(rotation)
    doc.save(path)
    return path


def png(path, size=(400, 300), color="navy"):
    Image.new("RGB", size, color).save(path)
    return path


@pytest.fixture
def d(tmp_path):
    (tmp_path / "out").mkdir()
    return tmp_path


def text(path):
    return "".join(p.get_text() for p in pymupdf.open(path))


def test_merge_split_rotate(d):
    a, b = make_pdf(d / "a.pdf", 2), make_pdf(d / "b.pdf", 3)
    merged = pages.merge([a, b], d / "out")
    assert len(pikepdf.open(merged).pages) == 5

    z = pages.split([merged], d / "out", mode="range", ranges="1-2, 3-5")
    assert len(zipfile.ZipFile(z).namelist()) == 2
    assert len(zipfile.ZipFile(pages.split([merged], d / "out", mode="every", every="2")).namelist()) == 3
    one = pages.split([merged], d / "out", mode="extract", ranges="2,4")
    assert len(pikepdf.open(one).pages) == 2
    with pytest.raises(UserError):
        pages.split([merged], d / "out", ranges="9")

    r = pages.rotate([a], d / "out", angle="90", pages="2")
    assert [p.rotation for p in pymupdf.open(r)] == [0, 90]


def test_organize(d):
    a, b = make_pdf(d / "a.pdf"), make_pdf(d / "b.pdf", 1, "Lain")
    spec = [{"f": 0, "p": 2, "r": 90}, {"f": None}, {"f": 1, "p": 0}, {"f": 0, "p": 2, "r": 0}]
    doc = pymupdf.open(pages.organize([a, b], d / "out", spec=json.dumps(spec)))
    assert len(doc) == 4
    assert "Halaman 3" in doc[0].get_text() and doc[0].rotation == 90
    assert doc[1].get_text().strip() == ""
    assert "Lain 1" in doc[2].get_text()


@pytest.mark.parametrize("rotation", [0, 90])
def test_crop(d, rotation):
    a = make_pdf(d / "a.pdf", rotation=rotation)
    area = json.dumps({"p": 0, "r": [0, 0, 0.5, 0.25]})
    doc = pymupdf.open(pages.crop([a], d / "out", area=area, scope="page"))
    full = pymupdf.open(a)[0].rect
    assert doc[0].rect.width == pytest.approx(full.width / 2, abs=1)
    assert doc[0].rect.height == pytest.approx(full.height / 4, abs=1)
    assert doc[1].rect == full


def test_numbers_and_watermark(d):
    a = make_pdf(d / "a.pdf")
    t = text(pages.page_numbers([a], d / "out", fmt="{n} / {total}", position="bottom-right"))
    assert "1 / 3" in t and "3 / 3" in t
    assert "RAHASIA" in text(pages.watermark([a], d / "out", text="RAHASIA"))
    img = png(d / "w.png")
    doc = pymupdf.open(pages.watermark([a], d / "out", image=img, size="50"))
    assert doc[0].get_images()


def test_images_and_scan(d):
    imgs = [png(d / "1.png"), png(d / "2.jpg", (300, 500))]
    doc = pymupdf.open(convert.images_to_pdf(imgs, d / "out", size="a4"))
    assert len(doc) == 2 and doc[0].rect.width > doc[0].rect.height  # landscape image -> landscape page

    photo = Image.new("RGB", (800, 1000), (40, 40, 40))
    ImageDraw.Draw(photo).polygon([(120, 100), (700, 140), (680, 900), (90, 860)], fill="white")
    photo.save(d / "photo.jpg")
    assert len(pymupdf.open(convert.scan_to_pdf([d / "photo.jpg"], d / "out", mode="bw"))) == 1


def test_pdf_to_formats(d):
    a = make_pdf(d / "a.pdf")
    from docx import Document
    from pptx import Presentation

    docx = convert.pdf_to_word([a], d / "out")
    assert "Halaman 1" in "\n".join(p.text for p in Document(docx).paragraphs)
    assert len(Presentation(convert.pdf_to_pptx([a], d / "out")).slides) == 3
    assert len(zipfile.ZipFile(convert.pdf_to_images([a], d / "out", fmt="png", dpi="72")).namelist()) == 3
    assert convert.pdf_to_images([a], d / "out", pages="2").suffix == ".jpg"
    assert "Halaman 1" in convert.pdf_to_markdown([a], d / "out").read_text(encoding="utf-8")
    with pytest.raises(UserError):
        convert.pdf_to_excel([a], d / "out")


def test_pdf_to_excel(d):
    doc = pymupdf.open()
    page = doc.new_page()
    for r in range(4):
        page.draw_line((72, 72 + r * 20), (372, 72 + r * 20))
    for c in range(4):
        page.draw_line((72 + c * 100, 72), (72 + c * 100, 132))
    for r in range(3):
        for c in range(3):
            page.insert_text((80 + c * 100, 87 + r * 20), f"r{r}c{c}", fontsize=9)
    doc.save(d / "t.pdf")
    from openpyxl import load_workbook

    ws = load_workbook(convert.pdf_to_excel([d / "t.pdf"], d / "out")).active
    assert ws["B2"].value == "r1c1"


def test_protect_unlock(d):
    a = make_pdf(d / "a.pdf")
    locked = secure.protect([a], d / "out", password="rahasia")
    with pytest.raises(pikepdf.PasswordError):
        pikepdf.open(locked)
    with pytest.raises(UserError):
        secure.unlock([locked], d / "out", password="salah")
    assert "Halaman 1" in text(secure.unlock([locked], d / "out", password="rahasia"))


@pytest.mark.parametrize("rotation", [0, 90])
def test_redact_removes_text(d, rotation):
    a = make_pdf(d / "a.pdf", 1, rotation=rotation)
    page = pymupdf.open(a)[0]
    word = page.search_for("SECRET")[0] * page.rotation_matrix  # displayed coordinates
    w, h = page.rect.width, page.rect.height
    r = [word.x0 / w - .01, word.y0 / h - .01, word.x1 / w + .01, word.y1 / h + .01]
    t = text(secure.redact([a], d / "out", areas=json.dumps([{"p": 0, "r": r}])))
    assert "SECRET" not in t and "Halaman 1" in t


def test_sign_metadata_flatten(d):
    a = make_pdf(d / "a.pdf")
    signed = secure.sign([a], d / "out", area=json.dumps({"p": 1, "r": [.6, .8, .9, .9]}), image=png(d / "s.png"))
    assert pymupdf.open(signed)[1].get_images()
    assert pymupdf.open(secure.metadata([a], d / "out", title="Skripsi")).metadata["title"] == "Skripsi"

    doc = pymupdf.open(a)
    w = pymupdf.Widget()
    w.field_name, w.field_type, w.rect, w.field_value = "nama", pymupdf.PDF_WIDGET_TYPE_TEXT, pymupdf.Rect(72, 200, 300, 220), "Budi"
    doc[0].add_widget(w)
    doc.save(d / "form.pdf")
    flat = pymupdf.open(secure.flatten([d / "form.pdf"], d / "out"))
    assert not list(flat[0].widgets()) and "Budi" in flat[0].get_text()


def test_compare(d):
    a = make_pdf(d / "a.pdf", 2)
    b = make_pdf(d / "b.pdf", 2, "Bagian")
    doc = pymupdf.open(analyze.compare([a, b], d / "out"))
    assert len(doc) == 2 and "perbedaan" in doc[0].get_text()


@needs("soffice")
def test_office_to_pdf(d):
    from docx import Document
    from openpyxl import Workbook

    doc = Document()
    doc.add_paragraph("Laporan praktikum")
    doc.save(d / "a.docx")
    assert "Laporan praktikum" in text(convert.office_to_pdf([d / "a.docx"], d / "out"))
    wb = Workbook()
    wb.active.append(list(range(40)))
    wb.create_sheet("Dua").append(["x"])
    wb.save(d / "b.xlsx")
    assert len(pymupdf.open(convert.excel_to_pdf([d / "b.xlsx"], d / "out"))) == 2
    assert len(pymupdf.open(convert.excel_to_pdf([d / "b.xlsx"], d / "out", sheets="2"))) == 1


@needs("gs")
def test_compress(d):
    a = make_pdf(d / "a.pdf")
    out = secure.compress([a], d / "out", level="high")
    assert out.stat().st_size <= a.stat().st_size


@needs("ocrmypdf")
@needs("tesseract")
def test_ocr_and_pdfa(d):
    page = pymupdf.open(make_pdf(d / "a.pdf", 1))[0]
    pix = page.get_pixmap(dpi=200)
    img = pymupdf.open()
    img.new_page().insert_image(img[0].rect, pixmap=pix)
    img.save(d / "scan.pdf")
    assert "Halaman" in text(secure.ocr([d / "scan.pdf"], d / "out", lang="ind+eng"))
    assert "pdfaid" in str(pikepdf.open(convert.pdf_to_pdfa([d / "scan.pdf"], d / "out")).open_metadata())


def test_api_flow(d):
    from fastapi.testclient import TestClient

    from app.main import app

    def wait(job):
        for _ in range(60):
            s = c.get(f"/jobs/{job}").json()
            if s["status"] != "running":
                return s
            time.sleep(0.5)

    a, b = make_pdf(d / "a.pdf", 2), make_pdf(d / "b.pdf", 1)
    with TestClient(app) as c:
        files = [("files", ("a.pdf", a.read_bytes())), ("files", ("b.pdf", b.read_bytes()))]
        job = c.post("/api/gabung", files=files).json()["job"]
        s = wait(job)
        assert s["status"] == "done" and s["name"] == "a-gabung.pdf", s
        r = c.get(s["download"])
        assert r.status_code == 200 and "a-gabung.pdf" in r.headers["content-disposition"]
        assert len(pymupdf.open(stream=r.content)) == 3
        assert c.get("/api/history").json()[0]["tool"] == "gabung"

        # Another browser (no cookie) can't see the job.
        assert TestClient(app).get(f"/jobs/{job}").status_code == 404

        bad = c.post("/api/putar", files=[("files", ("a.pdf", b"not a pdf"))])
        assert bad.status_code == 400
        wrong_ext = c.post("/api/putar", files=[("files", ("a.exe", a.read_bytes()))])
        assert wrong_ext.status_code == 400
        locked = secure.protect([a], d / "out", password="rahasia")
        r = c.post("/api/putar", files=[("files", ("x.pdf", locked.read_bytes()))])
        assert r.status_code == 400 and "terkunci" in r.json()["error"]
        tools = c.get("/api/tools").json()["tools"]
        assert len(tools) > 20 and all("fn" not in t and t["short"] for t in tools)
        cats = c.get("/api/tools").json()["categories"]
        assert all(t["category"] in cats for t in tools)

        # Batch: a single-file tool on two files returns a ZIP with one result per file.
        s = wait(c.post("/api/putar", files=files, data={"angle": "90"}).json()["job"])
        assert s["status"] == "done" and s["name"].endswith(".zip"), s
        names = zipfile.ZipFile(io.BytesIO(c.get(s["download"]).content)).namelist()
        assert sorted(names) == ["a-putar.pdf", "b-putar.pdf"]

        stats = c.get("/api/stats").json()["rows"]
        assert {r["tool"] for r in stats} >= {"gabung", "putar"}


def test_new_page_tools(d):
    a = make_pdf(d / "a.pdf", 5)
    left = pymupdf.open(pages.delete_pages([a], d / "out", pages="2,4-5"))
    assert [p.get_text().split("\n")[0] for p in left] == ["Halaman 1", "Halaman 3"]
    with pytest.raises(UserError):
        pages.delete_pages([a], d / "out", pages="1-5")

    two = pymupdf.open(stream=pages.nup([a], d / "out", per="2").read_bytes())  # Bytes: don't lock the file on Windows.
    assert len(two) == 3 and two[0].rect.width > two[0].rect.height
    assert len(pymupdf.open(stream=pages.nup([a], d / "out", per="4").read_bytes())) == 2

    f4 = pymupdf.open(pages.resize([a], d / "out", size="f4"))
    assert (round(f4[0].rect.width), round(f4[0].rect.height)) == (595, 936) and "Halaman 1" in f4[0].get_text()


def test_extract_and_html(d):
    a = make_pdf(d / "a.pdf", 2)
    assert "--- Halaman 2 ---" in convert.pdf_to_text([a], d / "out").read_text(encoding="utf-8")
    with pytest.raises(UserError):
        convert.extract_images([a], d / "out")
    doc = pymupdf.open(a)
    doc[0].insert_image(pymupdf.Rect(72, 200, 272, 350), filename=str(png(d / "p.png")))
    doc.save(d / "img.pdf")
    assert len(zipfile.ZipFile(convert.extract_images([d / "img.pdf"], d / "out")).namelist()) == 1

    secret = png(d / "secret.png")
    (d / "x.html").write_text(f"<h1>Tugas Akhir</h1><p>Isi</p><img src='{secret}'><img src='file:///{secret}'>",
                              encoding="utf-8")
    out = pymupdf.open(convert.html_to_pdf([d / "x.html"], d / "out"))
    assert "Tugas Akhir" in out[0].get_text()
    assert not out[0].get_images()  # Local files must not be readable through HTML.


def test_repair_and_form(d):
    a = make_pdf(d / "a.pdf", 2)
    broken = a.read_bytes()
    broken = broken[: broken.rfind(b"xref")]  # Chop off the cross-reference table and trailer.
    (d / "broken.pdf").write_bytes(broken)
    assert "Halaman 2" in text(secure.repair([d / "broken.pdf"], d / "out"))

    doc = pymupdf.open(a)
    for name, kind, y in [("nama", pymupdf.PDF_WIDGET_TYPE_TEXT, 200), ("setuju", pymupdf.PDF_WIDGET_TYPE_CHECKBOX, 240)]:
        w = pymupdf.Widget()
        w.field_name, w.field_type, w.rect = name, kind, pymupdf.Rect(72, y, 300, y + 20)
        doc[0].add_widget(w)
    doc.save(d / "form.pdf")
    filled = pymupdf.open(secure.fill_form([d / "form.pdf"], d / "out", fields=json.dumps({"nama": "Sari", "setuju": True})))
    values = {w.field_name: w.field_value for w in filled[0].widgets()}
    assert values["nama"] == "Sari" and values["setuju"] not in ("Off", False, "")
    with pytest.raises(UserError):
        secure.fill_form([a], d / "out", fields="{}")


def test_compare_visual_and_thesis(d):
    a = make_pdf(d / "a.pdf", 2)
    b = make_pdf(d / "b.pdf", 2, "Bagian")
    doc = pymupdf.open(analyze.compare([a, b], d / "out", mode="visual"))
    assert len(doc) == 2 and "2 halaman berbeda" in doc[0].get_text()

    # make_pdf puts text at x=72pt (2.5 cm) — inside a 4 cm left margin, so it must be flagged.
    report = pymupdf.open(stream=analyze.thesis_check([a], d / "out").read_bytes())
    first = report[0].get_text()
    assert len(report) == 3 and "di luar margin" in first and "semua halaman A4" in first
    ok = pymupdf.open(analyze.thesis_check([a], d / "out", left="2", top="1", right="1", bottom="1"))
    assert "semua isi berada di dalam margin" in ok[0].get_text()
