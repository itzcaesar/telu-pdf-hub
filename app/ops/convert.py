import io
import json
import zipfile

import pikepdf
import pymupdf
from PIL import Image, ImageOps

from . import UserError, parse_pages, run, soffice_profile

try:
    import pillow_heif
    pillow_heif.register_heif_opener()
except ImportError:  # HEIC support is optional outside Docker
    pass


def _soffice(src, out, filter_="pdf"):
    run(["soffice", soffice_profile(), "--headless", "--convert-to", filter_,
         "--outdir", str(out), str(src)], timeout=180)
    pdf = out / (src.stem + ".pdf")
    if not pdf.exists():
        raise UserError("Berkas tidak dapat dikonversi. Pastikan berkas tidak rusak atau terkunci.")
    return pdf


def office_to_pdf(files, out):
    return _soffice(files[0], out)


def excel_to_pdf(files, out, layout="fit", sheets=""):
    if layout != "fit":
        return _soffice(files[0], out)
    # One page per sheet, sized to fit the content, so columns never get cut.
    opts = json.dumps({"SinglePageSheets": {"type": "boolean", "value": "true"}})
    pdf = _soffice(files[0], out, "pdf:calc_pdf_Export:" + opts)
    if sheets.strip():
        src = pikepdf.open(pdf, allow_overwriting_input=True)
        keep = parse_pages(sheets, len(src.pages))
        dst = pikepdf.new()
        for i in keep:
            dst.pages.append(src.pages[i])
        dst.save(pdf)
    return pdf


PAGE_SIZES = {"a4": pymupdf.paper_size("a4"), "letter": pymupdf.paper_size("letter"),
              "f4": (595, 936)}


def images_to_pdf(files, out, size="a4", orientation="auto", margin="10"):
    doc = pymupdf.open()
    m = float(margin) * 72 / 25.4
    for f in files:
        im = ImageOps.exif_transpose(Image.open(f)).convert("RGB")
        _add_image_page(doc, im, size, orientation, m)
    path = out / "gambar.pdf"
    doc.save(path, garbage=3, deflate=True)
    return path


def _add_image_page(doc, im, size, orientation, m):
    if size == "fit":
        w, h = im.width * 0.75 + 2 * m, im.height * 0.75 + 2 * m  # 96 dpi
    else:
        w, h = PAGE_SIZES[size]
        landscape = im.width > im.height if orientation == "auto" else orientation == "landscape"
        if landscape:
            w, h = h, w
    page = doc.new_page(width=w, height=h)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=90)
    page.insert_image(pymupdf.Rect(m, m, w - m, h - m), stream=buf.getvalue(), keep_proportion=True)


def scan_to_pdf(files, out, mode="color"):
    import cv2
    import numpy as np

    doc = pymupdf.open()
    for f in files:
        im = np.array(ImageOps.exif_transpose(Image.open(f)).convert("RGB"))
        im = _deskew_document(cv2, np, im)
        if mode == "gray":
            im = cv2.cvtColor(im, cv2.COLOR_RGB2GRAY)
        elif mode == "bw":
            gray = cv2.cvtColor(im, cv2.COLOR_RGB2GRAY)
            im = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                                       cv2.THRESH_BINARY, 31, 15)
        _add_image_page(doc, Image.fromarray(im).convert("RGB"), "a4", "auto", 0)
    path = out / "scan.pdf"
    doc.save(path, garbage=3, deflate=True)
    return path


def _deskew_document(cv2, np, im):
    """Find the paper outline (largest 4-corner contour) and warp it flat. Returns input if none."""
    scale = 1000 / max(im.shape[:2])
    small = cv2.resize(im, None, fx=scale, fy=scale) if scale < 1 else im
    scale = min(scale, 1)
    gray = cv2.GaussianBlur(cv2.cvtColor(small, cv2.COLOR_RGB2GRAY), (5, 5), 0)
    edges = cv2.dilate(cv2.Canny(gray, 50, 150), np.ones((3, 3), np.uint8))
    contours, _ = cv2.findContours(edges, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)
    area = small.shape[0] * small.shape[1]
    for c in sorted(contours, key=cv2.contourArea, reverse=True)[:5]:
        approx = cv2.approxPolyDP(c, 0.02 * cv2.arcLength(c, True), True)
        if len(approx) == 4 and cv2.contourArea(approx) > 0.25 * area:
            pts = approx.reshape(4, 2).astype("float32") / scale
            s, d = pts.sum(1), np.diff(pts, axis=1).ravel()
            tl, br, tr, bl = pts[s.argmin()], pts[s.argmax()], pts[d.argmin()], pts[d.argmax()]
            w = int(max(np.linalg.norm(br - bl), np.linalg.norm(tr - tl)))
            h = int(max(np.linalg.norm(tr - br), np.linalg.norm(tl - bl)))
            dst = np.array([[0, 0], [w, 0], [w, h], [0, h]], dtype="float32")
            m = cv2.getPerspectiveTransform(np.array([tl, tr, br, bl]), dst)
            return cv2.warpPerspective(im, m, (w, h))
    return im


def pdf_to_word(files, out):
    from pdf2docx import Converter

    path = out / "hasil.docx"
    cv = Converter(str(files[0]))
    try:
        cv.convert(str(path))
    finally:
        cv.close()
    return path


def pdf_to_pptx(files, out, dpi="150"):
    from pptx import Presentation
    from pptx.util import Emu

    doc = pymupdf.open(files[0])
    prs = Presentation()
    first = doc[0].rect
    prs.slide_width, prs.slide_height = Emu(int(first.width * 12700)), Emu(int(first.height * 12700))
    for page in doc:
        png = page.get_pixmap(dpi=int(dpi)).tobytes("png")
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(png), 0, 0, prs.slide_width, prs.slide_height)
    path = out / "hasil.pptx"
    prs.save(path)
    return path


def pdf_to_excel(files, out):
    import pdfplumber
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    with pdfplumber.open(files[0]) as pdf:
        for p, page in enumerate(pdf.pages, 1):
            for t, table in enumerate(page.extract_tables(), 1):
                ws = wb.create_sheet(f"Hal {p} - Tabel {t}"[:31])
                for row in table:
                    ws.append([c if c is None else str(c) for c in row])
    if not wb.sheetnames:
        raise UserError("Tidak ada tabel yang terdeteksi pada PDF ini.")
    path = out / "hasil.xlsx"
    wb.save(path)
    return path


def pdf_to_images(files, out, fmt="jpg", dpi="150", pages=""):
    doc = pymupdf.open(files[0])
    idx = parse_pages(pages, len(doc))
    names = []
    for i in idx:
        name = out / f"halaman-{i + 1}.{fmt}"
        doc[i].get_pixmap(dpi=int(dpi)).save(name, jpg_quality=90) if fmt == "jpg" \
            else doc[i].get_pixmap(dpi=int(dpi)).save(name)
        names.append(name)
    if len(names) == 1:
        return names[0]
    path = out / "gambar.zip"
    with zipfile.ZipFile(path, "w") as z:
        for n in names:
            z.write(n, n.name)
            n.unlink()
    return path


def pdf_to_markdown(files, out):
    import pymupdf4llm

    path = out / "hasil.md"
    path.write_text(pymupdf4llm.to_markdown(str(files[0])), encoding="utf-8")
    return path


def pdf_to_pdfa(files, out):
    path = out / "pdfa.pdf"
    run(["ocrmypdf", "--skip-text", "--output-type", "pdfa-2", "--quiet",
         str(files[0]), str(path)], timeout=300)
    return path


def extract_images(files, out, min_size="64"):
    doc = pymupdf.open(files[0])
    seen, names = set(), []
    for p, page in enumerate(doc, 1):
        for img in page.get_images(full=True):
            xref = img[0]
            if xref in seen:
                continue
            seen.add(xref)
            info = doc.extract_image(xref)
            if not info or min(info["width"], info["height"]) < int(min_size):
                continue  # Skip icons, bullets and other tiny decorations.
            name = out / f"hal{p}-gambar{len(names) + 1}.{info['ext']}"
            name.write_bytes(info["image"])
            names.append(name)
    if not names:
        raise UserError("Tidak ada gambar yang dapat diekstrak dari PDF ini.")
    path = out / "gambar.zip"
    with zipfile.ZipFile(path, "w") as z:
        for n in names:
            z.write(n, n.name)
            n.unlink()
    return path


def pdf_to_text(files, out):
    doc = pymupdf.open(files[0])
    text = "\n\n".join(f"--- Halaman {i} ---\n{page.get_text(sort=True).strip()}" for i, page in enumerate(doc, 1))
    if not any(page.get_text().strip() for page in doc):
        raise UserError("PDF ini tidak memiliki teks. Jalankan OCR PDF terlebih dahulu.")
    path = out / "hasil.txt"
    path.write_text(text, encoding="utf-8")
    return path


def html_to_pdf(files, out, size="a4", margin="15"):
    # Story renders without network or local file access (archive=None), so uploaded HTML can't pull in
    # server files or make requests. External images simply don't render.
    html = files[0].read_bytes().decode("utf-8", errors="replace")
    story = pymupdf.Story(html=html)
    w, h = PAGE_SIZES[size] if size in PAGE_SIZES else PAGE_SIZES["a4"]
    m = float(margin) * 72 / 25.4
    page_rect = pymupdf.Rect(0, 0, w, h)
    path = out / "hasil.pdf"
    writer = pymupdf.DocumentWriter(str(path))
    more, pages = True, 0
    while more:
        dev = writer.begin_page(page_rect)
        more, _ = story.place(page_rect + (m, m, -m, -m))
        story.draw(dev)
        writer.end_page()
        pages += 1
        if pages > 500:
            break
    writer.close()
    return path

